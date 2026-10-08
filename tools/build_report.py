"""
在学院实验报告模板的基础上填写内容，生成完整实验报告。

用法：
    python tools/build_report.py <项目根目录>

输出：
    docs/实验一-实验报告-带键盘事件的计算器.docx
"""

import os
import shutil
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

SONG = "宋体"
HEI = "黑体"
MONO = "Consolas"

BODY_SIZE = 12          # 小四
CODE_SIZE = 9
CAPTION_SIZE = 10.5

HEADER_FILL = "1F3864"  # 表头深蓝
ROW_TINT = "F2F6FA"     # 隔行浅蓝灰


# --------------------------------------------------------------------------
# 基础排版工具
# --------------------------------------------------------------------------

def style_run(run, font=SONG, size=BODY_SIZE, bold=False, color=None):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = font
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), font)
    rfonts.set(qn("w:ascii"), font)
    rfonts.set(qn("w:hAnsi"), font)


def para(cell, text="", font=SONG, size=BODY_SIZE, bold=False, color=None,
         align=None, indent=True, before=0, after=6, line=1.5):
    p = cell.add_paragraph()
    fmt = p.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    if indent and font == SONG and size == BODY_SIZE:
        fmt.first_line_indent = Pt(BODY_SIZE * 2)
    if align is not None:
        p.alignment = align
    if text:
        style_run(p.add_run(text), font, size, bold, color)
    return p


def heading(cell, text, size=14):
    return para(cell, text, font=HEI, size=size, bold=True,
                indent=False, before=14, after=6, line=1.3)


def subheading(cell, text, size=12):
    return para(cell, text, font=HEI, size=size, bold=True,
                indent=False, before=8, after=4, line=1.3)


def code_block(cell, code):
    """等宽字体的代码片段（不加深色底纹，避免形成装饰性色块）。"""
    lines = code.strip("\n").split("\n")
    for i, line in enumerate(lines):
        p = cell.add_paragraph()
        fmt = p.paragraph_format
        fmt.space_before = Pt(4 if i == 0 else 0)
        fmt.space_after = Pt(4 if i == len(lines) - 1 else 0)
        fmt.line_spacing = 1.15
        fmt.left_indent = Pt(18)
        style_run(p.add_run(line if line else " "), MONO, CODE_SIZE)


def caption(cell, text):
    return para(cell, text, font=SONG, size=CAPTION_SIZE,
                indent=False, align=WD_ALIGN_PARAGRAPH.CENTER,
                before=2, after=10, line=1.2)


def image_row(cell, images, cap=None):
    """images: [(路径, 宽度英寸), ...]，同一行并排插入。"""
    p = cell.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    for idx, (path, width) in enumerate(images):
        if idx:
            style_run(p.add_run("    "), SONG, 10)
        p.add_run().add_picture(path, width=Inches(width))
    if cap:
        caption(cell, cap)


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def set_borders(table, color="D9D9D9", size=6):
    tblpr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + edge)
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(size))
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), color)
        borders.append(e)
    tblpr.append(borders)


def make_table(cell, headers, rows, widths, size=10.5):
    table = cell.add_table(rows=len(rows) + 1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_borders(table)

    # 表头行跨页时重复显示
    header_trpr = table.rows[0]._tr.get_or_add_trPr()
    header_trpr.append(OxmlElement("w:tblHeader"))

    def write(row_cells, values, is_header):
        for idx, (c, value) in enumerate(zip(row_cells, values)):
            c.width = Inches(widths[idx])
            p = c.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            p.alignment = (WD_ALIGN_PARAGRAPH.CENTER if is_header
                           else WD_ALIGN_PARAGRAPH.LEFT)
            style_run(p.add_run(value), SONG, size, bold=is_header,
                      color="FFFFFF" if is_header else None)
            if is_header:
                shade(c, HEADER_FILL)
            elif row_cells is not table.rows[1].cells:
                pass

    write(table.rows[0].cells, headers, True)
    for i, values in enumerate(rows):
        row = table.rows[i + 1]
        write(row.cells, values, False)
        if i % 2 == 1:
            for c in row.cells:
                shade(c, ROW_TINT)
    return table


# --------------------------------------------------------------------------
# 主流程
# --------------------------------------------------------------------------

def build(root):
    docs = os.path.join(root, "docs")
    shots = os.path.join(docs, "screenshots")
    template = os.path.join(docs, "实验报告模板.docx")
    out_path = os.path.join(docs, "实验一-实验报告-带键盘事件的计算器.docx")
    shutil.copyfile(template, out_path)

    doc = Document(out_path)
    table = doc.tables[0]

    # ---- 填写表头信息 ----
    def set_cell(cell, text):
        p = cell.paragraphs[0]
        for r in list(p.runs):
            r._element.getparent().remove(r._element)
        style_run(p.add_run(text), SONG, 12)

    set_cell(table.rows[1].cells[1], "王俊康")          # 姓名
    set_cell(table.rows[2].cells[3], "2026 年 10 月 8 日")  # 实验日期

    # ---- 报告正文写入模板最后一个单元格 ----
    body = table.rows[4].cells[0]

    # 确保长内容可以跨页
    trpr = table.rows[4]._tr.get_or_add_trPr()
    for cant in trpr.findall(qn("w:cantSplit")):
        trpr.remove(cant)

    para(body, "实验报告正文", font=HEI, size=16, bold=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, before=18, after=10, line=1.2)

    para(body, "本次实验用 Qt 6 与 Qt Widgets 完成了一个带键盘事件的计算器："
               "界面在 Qt 设计器中搭建，用垂直布局与网格布局安排显示屏和 19 个按键；"
               "运算部分用第一个操作数、待处理运算符和“是否等待新操作数”三个状态量管理，"
               "鼠标点击与键盘输入调用同一组处理函数。程序支持四则运算、小数、退格、清除，"
               "并处理了连续小数点、连续运算符、除数为 0、结果后继续输入等边界情况。"
               "全部 23 项自动化测试通过，代码通过 Git 分阶段提交。下面按实验要求的条目逐项说明。")

    # ---------- 一 ----------
    heading(body, "一、UI 界面布局与设计")
    para(body, "主界面在 Qt 设计器中完成，保存为 calculator.ui。窗口用 QVBoxLayout 作为顶层垂直布局，"
               "自上而下依次是显示运算过程的 QLabel（对象名 expressionLabel）、"
               "显示数据与结果的 QLineEdit（对象名 display）以及一个 QGridLayout（对象名 gridLayout）。"
               "网格布局是 4 列 5 行，共 19 个 QPushButton：第一行是 C、⌫、%、÷，"
               "第二到第四行是数字 7~9、4~6、1~3 及乘号、减号、加号，最后一行是占两列宽的数字 0、小数点和等号。"
               "窗口缩放时由布局自动分配空间，不需要手工计算控件坐标；"
               "显示屏设为只读并右对齐，运算过程显示在它上方的小字标签里。")
    para(body, "所有控件都显式命名，代码按对象名查找并连接，没有沿用 button、button_2 这类默认名称。"
               "主要控件与命名见表 1。")

    make_table(
        body,
        ["对象名", "控件类型", "作用"],
        [
            ["expressionLabel", "QLabel", "显示运算过程，如“12.5 ×”"],
            ["display", "QLineEdit", "主显示屏，只读、右对齐"],
            ["button0 ~ button9", "QPushButton", "数字 0 ~ 9"],
            ["buttonDot", "QPushButton", "小数点"],
            ["buttonAdd / buttonSub / buttonMul / buttonDiv", "QPushButton", "加、减、乘、除"],
            ["buttonEquals", "QPushButton", "等于"],
            ["buttonClear", "QPushButton", "清除（C）"],
            ["buttonBackspace", "QPushButton", "退格（⌫）"],
            ["buttonPercent", "QPushButton", "百分号（%）"],
        ],
        widths=[2.5, 1.3, 2.6],
    )
    caption(body, "表 1  主要控件与命名")

    para(body, "关键代码（calculator.ui 中按键网格的片段）：", indent=False, before=6, after=2)
    code_block(body, """
<layout class="QGridLayout" name="gridLayout">
  <item row="0" column="0"><widget class="QPushButton" name="buttonClear">...</widget></item>
  <item row="0" column="1"><widget class="QPushButton" name="buttonBackspace">...</widget></item>
  <item row="0" column="3"><widget class="QPushButton" name="buttonDiv">...</widget></item>
  <item row="4" column="0" colspan="2"><widget class="QPushButton" name="button0">...</widget></item>
</layout>
""")

    image_row(body, [(os.path.join(shots, "01_初始界面.png"), 2.45)],
              "图 1  程序初始界面（Qt 设计器布局，鼠标与键盘共用同一套逻辑）")

    # ---------- 二 ----------
    heading(body, "二、按键事件处理")
    para(body, "数字、小数点、退格和清除四个操作分别由 enterDigit()、enterDot()、backspace() 和 clearAll() 完成，"
               "按钮的 clicked 信号和键盘事件都调用这几个函数，因此鼠标与键盘的行为完全一致。"
               "输入数字时，如果显示屏还是 0 就直接替换，避免出现“0123”这样的前导零；"
               "退格每次删除最后一位，删空后自动回到 0 并重新处于等待输入状态；"
               "清除则把第一个操作数、运算符和输入状态一起复位。")
    subheading(body, "1. 避免同一操作数重复输入小数点")
    para(body, "小数点的问题在于，如果允许第二次输入，显示屏会出现“1..”，"
               "之后用 toDouble() 解析就会失败，数值变成 0，后续运算全部出错。"
               "处理办法是在输入前检查显示屏里是否已经包含小数点，只要已经存在就直接返回，不再追加。")
    code_block(body, """
QString text = ui->display->text();
if (text.contains(QLatin1Char('.')))
    return;                      // 已经有一个小数点，忽略这次输入
if (text.isEmpty())
    text = QStringLiteral("0");
setDisplayText(text + QLatin1Char('.'));
""")
    para(body, "另外，当显示屏刚被运算符或等号清空时按小数点，程序会自动补成“0.”，"
               "这样直接按小数点输入小数也能正常工作。")
    subheading(body, "2. 测试结果")
    para(body, "输入 12.34 后按一次退格，显示屏由 12.34 变为 12.3（图 2）；"
               "直接按小数点得到 0.5；连续按三次小数点再加数字 2，最终显示 1.2，"
               "程序没有出现解析失败或显示异常。")
    image_row(body, [(os.path.join(shots, "05_退格.png"), 2.45),
                     (os.path.join(shots, "02_小数输入.png"), 2.45)],
              "图 2  退格（12.34 → 12.3）与小数输入")

    # ---------- 三 ----------
    heading(body, "三、双操作数处理")
    para(body, "计算器在任何时刻只需要记住“第一个操作数、待处理的运算符、第二个操作数是否已经开始输入”，"
               "因此用三个成员变量就能完整描述状态，见表 2。")
    make_table(
        body,
        ["成员变量", "含义", "变化时机"],
        [
            ["m_accumulator", "第一个操作数，连续运算时保存中间结果", "按下运算符、得到结果时更新"],
            ["m_pendingOp", "尚未参与计算的运算符（None 表示没有）", "按下运算符时设置，计算完成后清空"],
            ["m_freshOperand", "下一次数字输入是否替换显示屏内容", "按下运算符或等号后置为 true"],
        ],
        widths=[1.7, 2.4, 2.3],
    )
    caption(body, "表 2  计算状态变量")
    para(body, "按下运算符时，如果此时还没有待处理运算符，就把当前显示的数值记为第一个操作数；"
               "如果已经有待处理运算符，说明用户输入了第二个操作数，先算出中间结果并回显，"
               "再把这个结果作为新的第一个操作数，从而实现连续运算。"
               "计算失败（例如除数为 0）时不改变状态，只提示错误。")
    code_block(body, """
void Calculator::applyOperator(Op op)
{
    const double value = currentValue();
    if (m_pendingOp == Op::None) {
        m_accumulator = value;                 // 记下第一个操作数
    } else {
        double result = 0.0;
        if (!compute(m_accumulator, value, m_pendingOp, result)) {
            showError(QStringLiteral("除数不能为 0"));
            return;
        }
        m_accumulator = result;                // 中间结果作为新的第一个操作数
        setDisplayText(formatNumber(result));
    }
    m_pendingOp = op;
    m_freshOperand = true;                     // 等待第二个操作数
}
""")
    para(body, "按等号时调用 evaluate()：把第一个操作数、运算符和第二个操作数交给 compute() 计算，"
               "把结果显示出来并清空运算符；同时把 m_freshOperand 置为 true，"
               "这样接下来输入的数字会作为新的操作数，而不是接在结果后面。"
               "如果按等号时并没有待处理的运算符，就把当前显示直接作为结果。")

    # ---------- 四 ----------
    heading(body, "四、连续操作与边界处理")
    para(body, "连续运算的处理方式是：每按一次运算符就把上一次的中间结果保存到 m_accumulator，"
               "因此 1 + 2 + 3 + 4 = 会依次得到 3、6、10。"
               "在实现过程中一共发现并修复了 5 个边界问题，记录见表 3。")

    make_table(
        body,
        ["测试输入", "原来的现象", "原因分析", "修改方法", "修改后结果"],
        [
            ["1 再连按三次小数点 2",
             "第一版允许继续追加小数点，显示屏出现 1..，按等号后解析失败变成 0",
             "enterDot() 没有检查显示屏中是否已有小数点",
             "输入前判断 text.contains('.')，已存在则直接返回",
             "显示屏保持 1.2，输入过程正常（图 2）"],
            ["5 + − × 3 =",
             "第一版把“−”当成第二个操作数，用 0 参与计算，结果错误",
             "applyOperator() 没有区分“刚按下运算符、还没输入第二个操作数”的情况",
             "增加 m_freshOperand 判断，连续按运算符时只替换 m_pendingOp",
             "按最后一个运算符计算，5 × 3 = 15"],
            ["5 ÷ 0 =",
             "第一版直接返回 inf，显示屏出现“inf”，之后运算全部失效",
             "compute() 没有判断除数是否为 0",
             "compute() 返回 bool，除数为 0 时返回 false；界面显示“错误”并提示原因",
             "显示“错误”和“除数不能为 0”（图 3），输入数字即可重新计算"],
            ["2 + 3 = 之后继续输入 4",
             "第一版把 4 接在结果 5 后面，显示屏变成 54",
             "得到结果后没有标记“下一个数字是新操作数”",
             "evaluate() 结束后置 m_freshOperand = true，并在 enterDigit() 中重置累计值",
             "输入 4 时显示屏直接变为 4，开始新的计算"],
            ["输入 123 的过程中连按退格",
             "第一版删空后显示屏变成空白",
             "backspace() 没有处理删空的情况",
             "删空或只剩负号时重置为 0，并回到等待新操作数状态",
             "显示 0，随后可以继续输入或运算"],
        ],
        widths=[1.25, 1.45, 1.35, 1.4, 1.35],
        size=9,
    )
    caption(body, "表 3  异常与边界输入测试记录")

    image_row(body, [(os.path.join(shots, "06_除零错误.png"), 2.45),
                     (os.path.join(shots, "07_连续运算.png"), 2.45)],
              "图 3  除数为 0 的错误提示与连续运算 1+2+3+4=10")

    # ---------- 五 ----------
    heading(body, "五、使用样式表改善界面")
    para(body, "界面样式放在 style.qss 中，通过 resources.qrc 打进可执行文件，"
               "程序启动时用 setStyleSheet() 加载。样式表按功能给按钮分组："
               "数字键使用深灰底配白字，运算符和等号使用橙色强调，"
               "C、⌫、% 使用浅灰底配深色字，显示屏使用更深的底色和更大的字号，"
               "并统一加了圆角。")
    code_block(body, """
QPushButton {
    background-color: #3a3d4a;  color: #ffffff;
    border: none;  border-radius: 12px;  font-size: 22px;
}
QPushButton#buttonAdd, QPushButton#buttonSub,
QPushButton#buttonMul, QPushButton#buttonDiv { background-color: #ff9f0a; font-size: 26px; }
QLineEdit#display {
    background-color: #2b2d38;  color: #ffffff;
    border-radius: 12px;  font-size: 34px;
}
""")
    para(body, "主要改进有三点：一是用颜色区分数字键、运算符键和功能键，"
               "用户扫一眼就能分清按键类别；二是把显示屏字号调大并与按键拉开间距，"
               "读结果更清楚；三是给所有按钮统一圆角和悬停、按下状态，"
               "点击时有反馈。样式全部按对象名选择，没有改动任何逻辑代码。")
    image_row(body, [(os.path.join(shots, "04_计算结果.png"), 2.45),
                     (os.path.join(shots, "03_运算过程.png"), 2.45)],
              "图 4  样式表生效后的界面：结果显示与运算过程")

    # ---------- 六 ----------
    heading(body, "六、键盘事件处理")
    para(body, "键盘输入在 keyPressEvent() 中统一处理，按键与功能的对应关系见表 4。"
               "处理时先判断事件文本是不是单个数字（这样主键盘区和小键盘区的数字都能识别），"
               "其余按键再用 switch 按 Qt::Key 分派。")
    make_table(
        body,
        ["键盘按键", "功能", "调用的函数"],
        [
            ["0 ~ 9", "输入数字", "enterDigit()"],
            [". （句点或小键盘）", "小数点", "enterDot()"],
            ["+  -  *（或 x）  /", "加、减、乘、除", "applyOperator()"],
            ["Enter 或 =", "等于", "evaluate()"],
            ["Backspace", "退格", "backspace()"],
            ["Esc / Delete / C", "清除", "clearAll()"],
            ["%", "百分号", "applyPercent()"],
        ],
        widths=[1.9, 2.0, 2.5],
    )
    caption(body, "表 4  键盘按键与功能对应表")

    para(body, "鼠标与键盘复用同一套处理逻辑：按钮的 clicked 信号和键盘分支都只调用上面那几个函数，"
               "没有为键盘单独写一份运算代码，所以两种输入的结果必然一致。"
               "为了让键盘事件能到达主窗口，所有按钮和显示屏都设置为不接收焦点，"
               "这样即使用鼠标点过按钮，焦点仍留在窗口上，键盘输入不会失效。")
    code_block(body, """
void Calculator::keyPressEvent(QKeyEvent *event)
{
    const QString text = event->text();
    if (text.size() == 1 && text.at(0).isDigit()) {   // 数字
        enterDigit(text);
        event->accept();
        return;
    }
    switch (event->key()) {
    case Qt::Key_Period:    enterDot();               break;   // .
    case Qt::Key_Plus:      applyOperator(Op::Add);   break;   // +
    case Qt::Key_Asterisk:  applyOperator(Op::Multiply); break;// *
    case Qt::Key_Slash:     applyOperator(Op::Divide); break;  // /
    case Qt::Key_Return:    evaluate();               break;   // Enter
    case Qt::Key_Backspace: backspace();              break;   // 退格
    case Qt::Key_Escape:    clearAll();               break;   // 清除
    default: QWidget::keyPressEvent(event); return;
    }
    event->accept();
}
""")
    para(body, "测试时分别用键盘输入 6 * 7 = 和鼠标点击 6、×、7、=，"
               "两种方式得到的都是 42（图 5）；"
               "把两者混用（鼠标点击数字、键盘按运算符、再鼠标点击数字）同样能得到正确结果，"
               "说明鼠标与键盘确实走的是同一套逻辑。")
    image_row(body, [(os.path.join(shots, "08_键盘输入.png"), 2.45),
                     (os.path.join(shots, "09_鼠标输入.png"), 2.45)],
              "图 5  键盘输入与鼠标点击得到相同结果 6 × 7 = 42")

    # ---------- 七 ----------
    heading(body, "七、Git 提交日志")
    para(body, "代码用 Git 分阶段提交，完整地体现了“搭框架 → 做界面 → 写功能 → 修问题 → 补测试 → 写文档”的过程，"
               "提交记录见表 5，提交历史截图见图 6。")
    make_table(
        body,
        ["序号", "提交说明", "主要内容"],
        [
            ["1", "chore: 初始化 Qt 计算器项目骨架", "qmake 与 CMake 工程、入口文件、.gitignore"],
            ["2", "feat(ui): 使用 UI 设计器完成界面布局与控件命名", "calculator.ui"],
            ["3", "feat: 添加界面样式表与资源文件", "style.qss、resources.qrc"],
            ["4", "feat: 实现四则运算、键盘输入与基本操作", "calculator.h / calculator.cpp 第一版"],
            ["5", "fix: 处理连续小数点、连续运算符与除零等异常输入", "边界处理与错误提示"],
            ["6", "test: 添加计算器行为自动化测试（23 个用例）", "tests/ 目录"],
            ["7", "docs: 添加运行截图生成工具与实验运行截图", "报告插图"],
            ["8", "fix(ui): 补充退格按钮，移除冗余的正负号按钮", "修复界面缺少退格键的问题"],
            ["9", "docs: 更新运行截图与演示视频（含退格按钮）", "重新生成素材"],
            ["10", "docs: 添加报告生成脚本与实验报告模板", "tools/build_report.py 与报告模板"],
        ],
        widths=[0.7, 3.1, 2.6],
        size=9.5,
    )
    caption(body, "表 5  Git 提交记录")
    image_row(body, [(os.path.join(docs, "git-history.png"), 6.2)],
              "图 6  Git 提交历史")
    para(body, "本报告与演示视频作为第 11 次提交（docs: 添加实验报告与演示视频）一并提交。")

    # ---------- 八 ----------
    heading(body, "八、AI 辅助开发记录")
    para(body, "下面记录一次比较有代表性的 AI 辅助开发过程：除数为 0 时的错误处理。")
    subheading(body, "1. 遇到的问题")
    para(body, "第一版程序没有判断除数是否为 0，按 5 ÷ 0 = 时显示屏直接出现“inf”，"
               "而且因为内部状态已经变成无效值，之后继续输入数字也算不对。")
    subheading(body, "2. 使用的提示词")
    para(body, "“我用 Qt 写了一个计算器，做除法时按 5 ÷ 0 = 显示屏出现 inf，"
               "而且之后继续输入也不正常。请问怎么判断除数是不是 0，"
               "并让程序在给出错误提示之后还能继续使用？”", indent=False)
    subheading(body, "3. AI 给出的方案")
    para(body, "AI 建议在计算函数里判断除数是否为 0，如果为 0 就返回失败而不要继续做除法；"
               "同时建议用一个成员变量标记错误状态，出错后先把状态复位再响应新的输入；"
               "还提醒比较浮点数不要用 == 0，而应该用 Qt::qFuzzyIsNull()。")
    subheading(body, "4. 实际运行结果")
    para(body, "按建议修改后，按 5 ÷ 0 = 会显示错误提示，之后输入 7 并按 + 3 = 能正常得到 10，"
               "qFuzzyIsNull() 的用法也是正确的，这一部分直接用在了最终代码里。")
    subheading(body, "5. 存在的问题")
    para(body, "AI 的第一版方案把整句提示“除数不能为 0”直接写进显示屏，"
               "而显示屏字号是 34 像素，这句话在窗口里显示不全；"
               "另外它建议在每一个输入函数开头都重复写一遍“如果出错就复位”，代码比较啰嗦。")
    subheading(body, "6. 我做的修改")
    para(body, "把长提示拆开：显示屏只显示“错误”两个字，完整原因显示在上方操作过程的小字标签里，"
               "既看清了提示又不挤占显示屏；复位逻辑只保留在输入数字的入口处统一处理，去掉了重复判断；"
               "最后补了自动化测试，用 5 ÷ 0 = 显示错误、再输入 7 + 3 = 得到 10 这两个断言确认修改有效。")

    # ---------- 九 ----------
    heading(body, "九、测试结果")
    para(body, "程序另写了 tests/test_calculator.cpp，通过发送真实的键盘事件和按钮点击事件驱动界面，"
               "再读取显示屏文本来判断结果，共 11 组、23 项断言，全部通过。用例与结果见表 6。")
    make_table(
        body,
        ["用例", "测试内容", "期望结果", "实际结果"],
        [
            ["1", "12 + 34、7 − 10、6 × 7、8 ÷ 2", "46、-3、42、4", "一致"],
            ["2", "1.5 + 2.25、0.1 + 0.2", "3.75、0.3", "一致（未出现浮点误差）"],
            ["3", "连按小数点 1...2、直接按小数点", "1.2、0.5", "一致"],
            ["4", "5 + − × 3 =", "15", "一致"],
            ["5", "123 退格、退到空、退格后重新计算", "12、0、10", "一致"],
            ["6", "1 + 2 + 3 + 4 =", "10", "一致"],
            ["7", "5 ÷ 0 =、出错后输入 7、7 + 3 =", "错误、7、10", "一致"],
            ["8", "2 + 3 = 后输入 40、40 + 2 =", "5、40、42", "一致"],
            ["9", "2 + 3 = 之后按 × 4 =", "20", "一致"],
            ["10", "鼠标 12.5 × 4 =、鼠标与键盘混合 9 + 3 =", "50、12", "一致"],
            ["11", "12.34 退格成 12.3 后加 1", "13.3", "一致"],
        ],
        widths=[0.55, 2.8, 1.8, 1.65],
        size=9.5,
    )
    caption(body, "表 6  自动化测试用例与结果（23 项断言全部通过）")

    para(body, "按实验要求还录制了 2～3 分钟的演示视频 docs/演示视频.mp4，"
               "依次演示了鼠标完成带小数的运算（12.5 × 4 = 50）、键盘计算（1.5 + 2.25 = 3.75）、"
               "退格（键盘与鼠标各演示一次）、清除后重新计算、除数为 0 的处理、"
               "连续运算（1 + 2 + 3 + 4 = 10）以及两种异常输入（连续小数点、连续运算符），"
               "视频结尾展示了本项目的 Git 提交历史。")

    # ---------- 十 ----------
    heading(body, "十、实验总结")
    subheading(body, "1. 出错分析与解决过程")
    para(body, "最花时间的问题不是运算本身，而是两处与状态相关的错误。"
               "第一处是键盘输入时有时完全没反应，原因是按钮默认会抢占焦点，"
               "用鼠标点过按钮之后焦点落在按钮上，键盘事件就不再送到主窗口；"
               "把显示屏和所有按钮的焦点策略设为 NoFocus、让窗口自己持有焦点之后问题就消失了。"
               "第二处是连续按运算符时程序会用错误的数据参与计算，"
               "后来发现是缺少“是否在等待第二个操作数”这个状态，补上 m_freshOperand 判断后解决。"
               "这两个问题让我意识到，界面程序里“状态”比“公式”更容易出错。")
    subheading(body, "2. 本次实验小结")
    para(body, "通过这次实验我熟悉了 Qt 设计器搭界面、信号与槽的连接方式、"
               "QVBoxLayout 与 QGridLayout 的配合使用、键盘事件的重写，以及用样式表美化界面。"
               "同时也体会到把运算逻辑集中到少数几个函数里的好处：鼠标和键盘只要调用同一组函数，"
               "行为就自然一致，出现问题时也只需要改一个地方。"
               "另外，用自动化测试代替手工反复点击之后，每次修改都能很快确认没有把原来的功能改坏。")
    subheading(body, "3. AI 辅助开发的体会")
    para(body, "AI 在这个实验里帮了不少忙：它能很快给出判断除数为 0 的写法、"
               "提醒浮点数比较要用 qFuzzyIsNull()，也能解释 Qt 布局的使用方式。"
               "但它给出的代码不能直接照抄使用，比如把整句错误提示塞进显示屏，"
               "在实际界面里会显示不全；它的方案里也有重复的复位判断。"
               "我的做法是先把它给的思路跑起来，看真实运行结果，"
               "再按界面尺寸和代码结构去调整，最后用测试用例确认修改确实生效。"
               "总的说来，AI 负责提供思路和初稿，验证、调整和取舍还是得自己做。")

    doc.save(out_path)
    return out_path


if __name__ == "__main__":
    project_root = sys.argv[1] if len(sys.argv) > 1 else os.getcwd()
    print(build(project_root))
